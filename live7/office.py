"""M14: office documents for 7 Live — docx / xlsx / pptx / pdf.

The user's wish: "make me a Word report / Excel sheet / slides / PDF, read
this file" — in English or Punjabi. No LibreOffice on this Mac: everything is
built with python-docx / openpyxl / python-pptx / fpdf2 and read back with
pypdf.

Safety rules (same as live7.mac):
  * every payload is validated BEFORE anything runs;
  * every tool catches ALL exceptions and returns {"ok": False, ...};
  * file paths go through live7.mac._check_path (inside home, no hidden
    component, never ~/Library); relative paths are tried against docs_dir;
  * new files never overwrite an existing one ("report 2.docx", ...);
  * reversible actions push an inverse onto the shared undo stack;
  * xlsx cells that look like formulas but contain a dangerous function are
    stored as plain STRINGS, never as formulas (CSV/DDE injection class).
"""

from __future__ import annotations

import logging
import re
import shutil
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from sentinel_orchestrator import Capability, RiskLevel, ToolRegistry, ToolSpec

from live7.mac import _check_path

MAX_TEXT = 20000          # doc_read text cap
MAX_SHEET_ROWS = 200      # doc_read rows per sheet
_NAME_MAX = 80

# Formula functions that turn a spreadsheet into an attack (exfiltration,
# DDE command execution). Such a cell is stored as a plain string instead.
_FORMULA_BLOCK = re.compile(
    r"WEBSERVICE|HYPERLINK|CALL|EXEC|REGISTER|\|", re.IGNORECASE)

_FONT_DIR = Path("/System/Library/Fonts/Supplemental")
_GURMUKHI_FONT = _FONT_DIR / "Gurmukhi.ttf"
_LATIN_FONT = _FONT_DIR / "Arial Unicode.ttf"


def _err(msg: str) -> Dict[str, Any]:
    return {"ok": False, "error": msg}


def _is_gurmukhi(text: str) -> bool:
    return any("਀" <= ch <= "੿" for ch in text)


# --------------------------------------------------------------------------- #
# markdown-lite helpers ("# " H1, "## " H2, "- " bullet, else paragraph)      #
# --------------------------------------------------------------------------- #

def _md_lines(text: Any) -> List[Tuple[str, str]]:
    if not isinstance(text, str):
        text = ""
    lines: List[Tuple[str, str]] = []
    for raw in text.splitlines():
        if raw.startswith("## "):
            lines.append(("h2", raw[3:].strip()))
        elif raw.startswith("# "):
            lines.append(("h1", raw[2:].strip()))
        elif raw.startswith("- "):
            lines.append(("bullet", raw[2:].strip()))
        else:
            lines.append(("p", raw))
    return lines


def _md_plain(text: Any) -> str:
    return "\n".join(body for _, body in _md_lines(text))


# --------------------------------------------------------------------------- #
# numbers / formulas for xlsx                                                  #
# --------------------------------------------------------------------------- #

def _looks_numeric(value: str) -> bool:
    return bool(re.fullmatch(r"-?\d+(\.\d+)?", value.strip()))


def _coerce_cell(value: Any) -> Any:
    """Strings that look like numbers become numbers; safe formulas keep the
    '=' prefix (openpyxl stores them as formulas); dangerous formulas are
    returned as plain strings with the '=' stripped so openpyxl keeps data_type
    's'."""
    if not isinstance(value, str):
        return value
    v = value.strip()
    if v.startswith("="):
        if _FORMULA_BLOCK.search(v):
            return v.lstrip("=")          # plain string cell, never a formula
        return v                          # openpyxl: '=' prefix => formula
    if _looks_numeric(v):
        try:
            return int(v)
        except ValueError:
            return float(v)
    return value


# --------------------------------------------------------------------------- #
# create                                                                       #
# --------------------------------------------------------------------------- #

def _clean_name(raw: Any) -> Tuple[Optional[str], Optional[str]]:
    if not isinstance(raw, str):
        return None, "name must be a string"
    name = raw.replace("\\", "/").rsplit("/", 1)[-1]          # basename only
    name = name.strip().lstrip(".").strip()[:_NAME_MAX].rstrip(". ").strip()
    if not name:
        return None, "name must not be empty"
    return name, None


def _fresh_path(docs_dir: Path, name: str, kind: str) -> Path:
    path = docs_dir / f"{name}.{kind}"
    n = 1
    while path.exists():
        n += 1
        path = docs_dir / f"{name} {n}.{kind}"
    return path


def _write_docx(path: Path, text: Any) -> None:
    import docx
    d = docx.Document()
    for kind, body in _md_lines(text):
        if kind == "h1":
            d.add_heading(body, level=1)
        elif kind == "h2":
            d.add_heading(body, level=2)
        elif kind == "bullet":
            d.add_paragraph(body, style="List Bullet")
        else:
            d.add_paragraph(body)
    d.save(str(path))


def _write_xlsx(path: Path, rows: Any) -> Optional[str]:
    import openpyxl
    if rows is None:
        rows = []
    if not isinstance(rows, list):
        return "rows must be a list of rows"
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, row in enumerate(rows, start=1):
        cells = row if isinstance(row, list) else [row]
        for c, value in enumerate(cells, start=1):
            cell = ws.cell(row=r, column=c, value=_coerce_cell(value))
            if r == 1:
                cell.font = openpyxl.styles.Font(bold=True)
    wb.save(str(path))
    return None


def _write_pptx(path: Path, slides: Any) -> Optional[str]:
    from pptx import Presentation
    if slides is None:
        slides = []
    if not isinstance(slides, list):
        return "slides must be a list of {title, bullets}"
    prs = Presentation()
    for item in slides:
        if not isinstance(item, dict):
            return "each slide must be an object {title, bullets}"
        title = str(item.get("title") or "")
        bullets = item.get("bullets")
        bullets = [str(b) for b in bullets] if isinstance(bullets, list) else []
        # A slide with no bullets is a title slide.
        layout = prs.slide_layouts[0 if not bullets else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        if bullets:
            body = slide.placeholders[1].text_frame
            body.text = bullets[0]
            for b in bullets[1:]:
                body.add_paragraph().text = b
    prs.save(str(path))
    return None


logging.getLogger("fontTools").setLevel(logging.ERROR)


def _pdf_document(text: Any, heading_sizes: bool = True):
    """fpdf2 document from markdown-lite text. Gurmukhi lines use the Gurmukhi
    font, everything else Arial Unicode, always with text shaping on; every
    line restarts at the left margin."""
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    pdf.set_text_shaping(True)
    pdf.add_font("latin", "", str(_LATIN_FONT))
    pdf.add_font("gurmukhi", "", str(_GURMUKHI_FONT))
    # Gurmukhi.ttf has no Latin letters: on mixed lines ("7 Live — ਟੈਸਟ") they vanished.
    # Latin glyphs missing from the current font fall back to Arial Unicode.
    pdf.set_fallback_fonts(["latin"])
    for kind, body in _md_lines(text):
        font = "gurmukhi" if _is_gurmukhi(body) else "latin"
        size = 12
        if heading_sizes:
            size = {"h1": 18, "h2": 15, "bullet": 12, "p": 12}[kind]
        pdf.set_font(font, size=size)
        prefix = "- " if kind == "bullet" else ""
        pdf.multi_cell(0, 8, prefix + body, new_x="LMARGIN", new_y="NEXT")
    return pdf


def _write_pdf(path: Path, text: Any) -> None:
    _pdf_document(text).output(str(path))


# --------------------------------------------------------------------------- #
# read                                                                         #
# --------------------------------------------------------------------------- #

def _read_docx(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs if p.text is not None)


def _read_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    lines: List[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        bullets: List[str] = []
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        for shape in slide.shapes:
            if shape is slide.shapes.title or not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                t = "".join(run.text for run in para.runs).strip()
                if t:
                    bullets.append(t)
        lines.append(f"Slide {i}: {title}".rstrip())
        lines.extend(f"- {b}" for b in bullets)
    return "\n".join(lines)


def _read_pdf(path: Path) -> str:
    import pypdf
    reader = pypdf.PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_xlsx(path: Path) -> Dict[str, List[List[Any]]]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True)
    sheets: Dict[str, List[List[Any]]] = {}
    for ws in wb.worksheets:
        rows: List[List[Any]] = []
        for row in ws.iter_rows(max_row=MAX_SHEET_ROWS, values_only=True):
            rows.append(["" if v is None else v for v in row])
        sheets[ws.title] = rows
    wb.close()
    return sheets


# --------------------------------------------------------------------------- #
# edit                                                                         #
# --------------------------------------------------------------------------- #

def _edit_docx(path: Path, action: str, payload: Dict[str, Any]) -> Optional[str]:
    import docx
    d = docx.Document(str(path))
    if action == "append":
        for kind, body in _md_lines(payload.get("text")):
            if kind == "h1":
                d.add_heading(body, level=1)
            elif kind == "h2":
                d.add_heading(body, level=2)
            elif kind == "bullet":
                d.add_paragraph(body, style="List Bullet")
            else:
                d.add_paragraph(body)
    elif action == "replace":
        find = payload.get("find")
        if not isinstance(find, str) or not find:
            return "replace needs a non-empty 'find' string"
        text = payload.get("text")
        text = text if isinstance(text, str) else ""
        found = False
        for para in d.paragraphs:
            if find in para.text:
                found = True
                for run in para.runs:
                    if find in run.text:
                        run.text = run.text.replace(find, text)
        if not found:
            return f"not found: {find!r}"
    else:
        return f"action {action!r} not supported for docx"
    d.save(str(path))
    return None


def _edit_xlsx(path: Path, action: str, payload: Dict[str, Any]) -> Optional[str]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path))
    ws = wb.worksheets[0]
    if action == "set_cell":
        ref = payload.get("cell")
        if not isinstance(ref, str) or not re.fullmatch(r"[A-Za-z]{1,2}\d{1,4}",
                                                        ref.strip()):
            wb.close()
            return "cell must be an A1-style reference like B2"
        ws[ref.strip().upper()] = _coerce_cell(payload.get("value"))
    elif action == "append_row":
        row = payload.get("row")
        if not isinstance(row, list):
            wb.close()
            return "append_row needs a 'row' list"
        ws.append([_coerce_cell(v) for v in row])
    else:
        wb.close()
        return f"action {action!r} not supported for xlsx"
    wb.save(str(path))
    wb.close()
    return None


def _edit_pptx(path: Path, action: str, payload: Dict[str, Any]) -> Optional[str]:
    from pptx import Presentation
    if action != "add_slide":
        return f"action {action!r} not supported for pptx"
    slide = payload.get("slide")
    if not isinstance(slide, dict):
        return "add_slide needs a 'slide' object {title, bullets}"
    prs = Presentation(str(path))
    title = str(slide.get("title") or "")
    bullets = slide.get("bullets")
    bullets = [str(b) for b in bullets] if isinstance(bullets, list) else []
    layout = prs.slide_layouts[0 if not bullets else 1]
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = title
    if bullets:
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
    prs.save(str(path))
    return None


# --------------------------------------------------------------------------- #
# registration                                                                 #
# --------------------------------------------------------------------------- #

def register_office_tools(registry: ToolRegistry, undo=None,
                          docs_dir: Optional[Path] = None,
                          home: Optional[Path] = None,
                          show=None) -> None:
    """show: optional async callable(path) that brings a finished document to the front
    (live7.front.show_front in the app; None in tests). A payload "show": false skips it."""
    home = Path(home) if home is not None else Path.home()
    docs_dir = Path(docs_dir) if docs_dir is not None else home / "Documents" / "7"

    async def shown(result: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        if show is not None and result.get("ok") and payload.get("show", True) is not False:
            try:
                result["shown"] = bool(await show(result["path"]))
            except Exception:
                result["shown"] = False
        return result

    def push_undo(label: str, undo_fn) -> None:
        if undo is not None:
            undo.push(label, undo_fn)

    def resolve(raw: Any) -> Tuple[Optional[Path], Optional[str]]:
        """Relative paths are tried against docs_dir first; everything still
        goes through live7.mac._check_path."""
        if isinstance(raw, str) and raw.strip():
            candidate = (docs_dir / raw.strip()).resolve()
            try:
                candidate.relative_to(home.resolve())
            except (OSError, RuntimeError, ValueError):
                candidate = None
            if candidate is not None and candidate.exists():
                return _check_path(str(candidate), home)
        return _check_path(raw, home)

    async def doc_create(payload: Dict[str, Any]) -> Dict[str, Any]:
        try:
            if not isinstance(payload, dict):
                return _err("payload must be an object")
            kind = payload.get("kind")
            if kind not in ("docx", "xlsx", "pptx", "pdf"):
                return _err("kind must be one of docx|xlsx|pptx|pdf")
            name, err = _clean_name(payload.get("name"))
            if err:
                return _err(err)
            docs_dir.mkdir(parents=True, exist_ok=True)
            path = _fresh_path(docs_dir, name, kind)
            if kind == "docx":
                _write_docx(path, payload.get("text"))
            elif kind == "pdf":
                _write_pdf(path, payload.get("text"))
            elif kind == "xlsx":
                err = _write_xlsx(path, payload.get("rows"))
                if err:
                    return _err(err)
            else:
                err = _write_pptx(path, payload.get("slides"))
                if err:
                    return _err(err)

            async def _undo_create(p=path):
                if p.exists():
                    p.unlink()
                return {"deleted": str(p)}
            push_undo(f"doc_create {path.name}", _undo_create)
            return await shown({"ok": True, "path": str(path)}, payload)
        except Exception as exc:
            return _err(str(exc))

    async def doc_read(payload: Dict[str, Any]) -> Dict[str, Any]:
        try:
            if not isinstance(payload, dict):
                return _err("payload must be an object")
            path, err = resolve(payload.get("path"))
            if err:
                return _err(err)
            if not path.is_file():
                return _err(f"no such file: {payload.get('path')}")
            kind = path.suffix.lower().lstrip(".")
            if kind not in ("docx", "xlsx", "pptx", "pdf", "txt", "md", "csv"):
                return _err(f"unsupported document type: .{kind}")
            out: Dict[str, Any] = {"ok": True, "kind": kind}
            if kind == "docx":
                text = _read_docx(path)
            elif kind == "pptx":
                text = _read_pptx(path)
            elif kind == "pdf":
                text = _read_pdf(path)
            elif kind == "xlsx":
                sheets = _read_xlsx(path)
                out["sheets"] = sheets
                text = "\n".join(
                    f"{name}: {len(rows)} rows" for name, rows in sheets.items())
            else:
                text = path.read_text(encoding="utf-8", errors="replace")
            out["text"] = text[:MAX_TEXT]
            if len(text) > MAX_TEXT:
                out["truncated"] = True
            return out
        except Exception as exc:
            return _err(str(exc))

    async def doc_edit(payload: Dict[str, Any]) -> Dict[str, Any]:
        try:
            if not isinstance(payload, dict):
                return _err("payload must be an object")
            path, err = resolve(payload.get("path"))
            if err:
                return _err(err)
            if not path.is_file():
                return _err(f"no such file: {payload.get('path')}")
            kind = path.suffix.lower().lstrip(".")
            if kind == "pdf":
                return _err("pdf is not editable — create a new one with "
                            "doc_create instead")
            if kind not in ("docx", "xlsx", "pptx", "txt", "md", "csv"):
                return _err(f"unsupported document type: .{kind}")
            action = payload.get("action")
            if not isinstance(action, str):
                return _err("action must be a string")

            # Backups live in home/.seven/doc_backups — never as clutter next to the user's
            # file; unique per edit so consecutive edits keep separate undo points.
            backup_dir = home / ".seven" / "doc_backups"
            backup_dir.mkdir(parents=True, exist_ok=True)
            backup = backup_dir / f"{time.time_ns()}-{path.name}"
            shutil.copyfile(path, backup)   # content bytes; never metadata
            try:
                if kind == "docx":
                    err = _edit_docx(path, action, payload)
                elif kind == "xlsx":
                    err = _edit_xlsx(path, action, payload)
                elif kind == "pptx":
                    err = _edit_pptx(path, action, payload)
                else:  # txt / md / csv
                    if action == "append":
                        text = payload.get("text")
                        text = text if isinstance(text, str) else ""
                        with path.open("a", encoding="utf-8") as fh:
                            fh.write(text + "\n")
                        err = None
                    elif action == "replace":
                        find = payload.get("find")
                        if not isinstance(find, str) or not find:
                            err = "replace needs a non-empty 'find' string"
                        else:
                            body = path.read_text(encoding="utf-8")
                            if find not in body:
                                err = f"not found: {find!r}"
                            else:
                                repl = payload.get("text")
                                repl = repl if isinstance(repl, str) else ""
                                path.write_text(body.replace(find, repl),
                                                encoding="utf-8")
                                err = None
                    else:
                        err = f"action {action!r} not supported for {kind}"
            except Exception:
                shutil.move(str(backup), str(path))   # restore on any failure
                raise
            if err:
                shutil.move(str(backup), str(path))   # no change made
                return _err(err)

            async def _undo_edit(p=path, b=backup):
                shutil.move(str(b), str(p))
                return {"restored": str(p)}
            push_undo(f"doc_edit {path.name}", _undo_edit)
            return await shown({"ok": True, "path": str(path)}, payload)
        except Exception as exc:
            return _err(str(exc))

    async def doc_to_pdf(payload: Dict[str, Any]) -> Dict[str, Any]:
        try:
            if not isinstance(payload, dict):
                return _err("payload must be an object")
            path, err = resolve(payload.get("path"))
            if err:
                return _err(err)
            if not path.is_file():
                return _err(f"no such file: {payload.get('path')}")
            kind = path.suffix.lower().lstrip(".")
            if kind not in ("docx", "pptx", "md", "txt"):
                return _err("doc_to_pdf converts docx, pptx, md or txt")
            if kind == "docx":
                text = _read_docx(path)
            elif kind == "pptx":
                text = _read_pptx(path)
            else:
                text = path.read_text(encoding="utf-8", errors="replace")
            docs_dir.mkdir(parents=True, exist_ok=True)
            out_path = _fresh_path(docs_dir, path.stem, "pdf")
            _write_pdf(out_path, text)

            async def _undo_pdf(p=out_path):
                if p.exists():
                    p.unlink()
                return {"deleted": str(p)}
            push_undo(f"doc_to_pdf {out_path.name}", _undo_pdf)
            return await shown({"ok": True, "path": str(out_path),
                                "note": "text-only conversion (no images or layout)"}, payload)
        except Exception as exc:
            return _err(str(exc))

    L, M = RiskLevel.LOW, RiskLevel.MEDIUM
    NONE = Capability.NONE
    SENS = Capability.SENSITIVE_ACCESS
    EXT = Capability.EXTERNAL_EFFECT
    for spec in (
        ToolSpec("doc_create", doc_create, M, NONE),
        ToolSpec("doc_read", doc_read, L, SENS),
        ToolSpec("doc_edit", doc_edit, M, EXT),
        ToolSpec("doc_to_pdf", doc_to_pdf, M, NONE),
    ):
        registry.register(spec)
